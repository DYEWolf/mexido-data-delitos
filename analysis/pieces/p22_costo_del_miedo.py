"""Pieza 22 — ¿Qué deja de hacer la gente en Jalisco por miedo, y cómo se protege?

Datos: ENVIPE 2026 (febrero-abril de 2026), réplica con ENVIPE 2025; tper_vic1, resto del módulo AP4. Variables
verificadas en el diccionario y los catálogos de cada ZIP antes de calcular (25-09-2026):
  - AP4_10_01..16: dejó de hacer la actividad por temor a ser víctima (1 sí, 2 no, 3 no aplica, 9 no sabe / no
    responde); las 16 existen con la misma etiqueta en las dos ediciones. AP4_10_17 (llegar muy tarde o dejar la casa
    sola) y AP4_10A_1..4 (actividades en línea) solo existen en 2026.
  - AP4_11_01..11: medidas de protección contra la delincuencia (1 sí, 2 no, 9 no sabe / no responde).
  - AP4_12: gasto en protección, en pesos (0000000 no considera gasto, 9999999 no sabe / no responde, "b" blanco).
    Nota del 25-09-2026, al correr la pieza: en 2026 la pregunta también se hace a quien sólo tomó medidas
    electrónicas (AP4_11A, nuevas); el gasto no se compara entre ediciones. No afecta ninguna hipótesis.
  - AP4_8_1..6 problema en su colonia (1 sí, 2 no, 3 no aplica, 9) y AP4_9_1..6 los vecinos se organizan para
    resolverlo (1 sí, 2 no, 9; se pregunta a quien reporta el problema).
  - AP4_6_1..3: sensación de poder ser víctima de robo en la calle, lesiones o extorsión y secuestro.
  - SEXO, CVE_ENT, CVE_MUN, FAC_ELE, EST_DIS, UPM_DIS.

Transparencia. Ningún resultado de estas variables se había visto. Se conocen los de la pieza 16 (percepción: 78.9% se
siente inseguro en el estado; sin cambio 2025-2026, razón 1.02 [0.98-1.06]; más inseguridad en la colonia en el área
metropolitana, +23.5 puntos). Al revisar la estructura del tabulado de INEGI 5.32 (actividades cotidianas) se vio que
sus filas están ordenadas por la estimación: se conoce el orden nacional de las actividades (primero, permitir que los
menores salgan solos; después, salir de noche) y el de Aguascalientes, no sus cifras. Al depurar el lector del cuadro
(después de escribir las hipótesis y antes de ver cualquier cifra) también se vio el orden de Jalisco 2026: primero,
permitir que los menores salgan solos; después, usar joyas y salir de noche. Por eso no hay hipótesis sobre qué
actividad es la primera, y ninguna hipótesis depende de ese orden.

Definiciones. "Dejó de hacer" una actividad: proporción de 1 entre quienes la actividad les aplica (se excluye el código
3; "no sabe" queda en el denominador), que es como INEGI construye el cuadro 5.32 según su nota. "Al menos una": alguna
de las 16 actividades comparables en 1, entre todos los adultos. "Alguna medida": alguno de AP4_11_01..11 en 1, entre
todos los adultos. Organización vecinal: AP4_9_x = 1 entre quienes reportan el problema (AP4_8_x = 1).

Hipótesis (declaradas el 25-09-2026 antes de calcular):
  H22a. La mayoría de los adultos de Jalisco dejó de hacer al menos una de las 16 actividades por temor (2026):
        IC95 inferior > 50%. Refutación: IC95 inferior <= 50%.
  H22b. Dejar de hacer cosas por miedo no siguió la baja del homicidio registrado: la razón 2026/2025 de "al menos una"
        en Jalisco tiene IC95 inferior > 0.90 (como H16d). Refutación: IC95 inferior <= 0.90.
  H22c. En Jalisco, más mujeres que hombres dejaron de salir de noche (2026): IC95 inferior de la diferencia mujeres −
        hombres > 0. Refutación: IC95 inferior <= 0.
  H22d. En el área metropolitana más adultos dejaron de hacer al menos una actividad que en el resto del estado (2026):
        IC95 inferior de la diferencia > 0. Refutación: IC95 inferior <= 0.
  H22e. En Jalisco, menos adultos toman alguna medida de protección que los que dejan de hacer al menos una actividad
        (2026, mismas personas): IC95 superior de la diferencia medidas − actividades < 0. Refutación: IC95 superior >= 0.
  H22f. Los vecinos se organizan menos contra el delito que por los servicios: en Jalisco (2026), entre quienes
        reportan el problema, la proporción que dice que los vecinos se organizan es menor para los robos que para la
        falta de alumbrado público: IC95 superior de la diferencia robos − alumbrado < 0. Refutación: IC95 superior >= 0.
Validación contra INEGI (declarada antes de leer cualquier cifra del cuadro 5.32; como H16h e H16i):
  H22g. Las estimaciones de "sí dejó de hacerla" reproducen el cuadro 5.32 en las 16 actividades comparables, Jalisco y
        nacional, en las dos ediciones (64 cifras): diferencia absoluta <= 0.5 puntos en todas.
  H22h. Los errores estándar de esas 64 cifras están entre 0.90 y 1.10 veces los publicados.

Descripción (sin hipótesis): cada actividad en Jalisco y en el país, en las dos ediciones; AP4_10_17 y las actividades
en línea (2026); cada medida de protección; gasto en protección (proporción que gasta, gasto medio entre quienes gastan
y total estatal, con IC95; "no sabe" fuera del promedio); los seis problemas de la colonia y la organización vecinal;
AP4_6; todo en Jalisco, área metropolitana, resto del estado y país.

Exploratorio (agregado el 25-09-2026 después de ver los resultados):
  E22i. Cambio 2025 -> 2026 de cada actividad en Jalisco y en el país (razón, como H22b). Las diferencias Jalisco −
        nacional por actividad y en "al menos una" y "alguna medida" se agregaron en el mismo momento, como descripción.

Estimador: el de las piezas 9, 13 y 16 (razón ponderada con FAC_ELE, linealización de Taylor con EST_DIS y UPM_DIS,
IC95 en escala logit para proporciones). Diferencias dentro de una edición: linealización sobre la muestra completa.
Entre ediciones: muestras independientes, varianzas sumadas, razón por método delta en escala logarítmica.
"""
from __future__ import annotations

import re
import zipfile

import numpy as np
import pandas as pd

from mvj import data as D
from pieces.p9_cifra_negra import ratio_se
from pieces.p13_no_denuncia_confianza import _catalog, _lin, _var, diff, load_personas, prop
from pieces.p16_percepcion import _between, _clean, _paired, _se

ACT = {"01": "Salir de noche", "02": "Permitir que los (las) menores de edad que viven en el hogar salgan solos(as)",
       "03": "Visitar parientes o amigos(as)", "04": "Tomar taxi", "05": "Usar transporte público",
       "06": "Llevar dinero en efectivo", "07": "Ir a la escuela", "08": "Ir al cine o al teatro", "09": "Salir a caminar",
       "10": "Usar joyas", "11": "Salir a comer o cenar", "12": "Llevar tarjeta de crédito o débito", "13": "Ir al estadio",
       "14": "Frecuentar centros comerciales", "15": "Viajar por carretera a otro estado o municipio",
       "16": "Llevar teléfono móvil o celular"}
# Palabra clave del diccionario de cada ZIP que debe aparecer en la etiqueta de cada código.
ACT_DIC = {"01": "salir de noche", "02": "menores", "03": "parientes", "04": "taxi", "05": "transporte público",
           "06": "efectivo", "07": "escuela", "08": "cine", "09": "caminar", "10": "joyas", "11": "comer",
           "12": "tarjeta", "13": "estadio", "14": "centros comerciales", "15": "carretera", "16": "móvil"}
MED = {"01": "Cambiar o reforzar puertas o ventanas", "02": "Cambiar, colocar o reforzar cerraduras o candados",
       "03": "Colocar o reforzar rejas o bardas", "04": "Instalar alarmas o videocámaras",
       "05": "Contratar vigilancia privada en la calle o colonia", "06": "Acciones conjuntas con vecinos",
       "07": "Contratar seguros", "08": "Comprar un perro guardián", "09": "Adquirir armas de fuego",
       "10": "Cambiarse de vivienda", "11": "Otra medida"}
PROB = {"1": "Falta de alumbrado público", "2": "Falta de agua", "3": "Baches o fugas de agua",
        "4": "Pandillerismo violento", "5": "Robos", "6": "Delincuencia en alrededores de escuelas"}
SENS = {"1": "Robo o asalto en la calle o transporte público", "2": "Lesiones por agresión física",
        "3": "Extorsión o secuestro"}
EN_LINEA = {"1": "Comprar o vender por internet", "2": "Banca electrónica", "3": "Comunicarse por internet",
            "4": "Otra actividad en línea"}


def _code(df, v) -> pd.Series:
    return df[v].fillna("").astype(str).str.strip()


def _dic(year: int) -> str:
    z = zipfile.ZipFile(D.PATHS["envipe"] / f"conjunto_de_datos_envipe_{year}_csv.zip")
    base = f"tper_vic1_envipe{year}"
    return z.read(f"{base}/diccionario_de_datos/diccionario_de_datos_{base}.csv").decode("utf-8", errors="replace").lower()


def _checks(df: pd.DataFrame, year: int) -> None:
    dic = _dic(year)
    for c, k in ACT_DIC.items():
        D.check(re.search(rf'"[^"]*{re.escape(k)}[^"]*","ap4_10_{c}"', dic) is not None, f"diccionario AP4_10_{c} {year}")
        D.check(set(_code(df, f"AP4_10_{c}")) <= {"1", "2", "3", "9", ""}, f"AP4_10_{c} {year}: códigos")
    for c in MED:
        D.check(set(_code(df, f"AP4_11_{c}")) <= {"1", "2", "9", ""}, f"AP4_11_{c} {year}: códigos")
    for c in PROB:
        D.check(set(_code(df, f"AP4_8_{c}")) <= {"1", "2", "3", "9", ""}, f"AP4_8_{c} {year}: códigos")
        a9 = _code(df, f"AP4_9_{c}")
        D.check(set(a9) <= {"1", "2", "9", "", "b"}, f"AP4_9_{c} {year}: códigos")
        D.check(((a9.isin(["1", "2", "9"])) == (_code(df, f"AP4_8_{c}") == "1")).all(),
                f"AP4_9_{c} {year}: se responde sólo y siempre cuando AP4_8_{c} = 1")
    g = _code(df, "AP4_12")
    D.check(g.str.fullmatch(r"\d{1,7}|b|").all(), f"AP4_12 {year}: formato")
    z = zipfile.ZipFile(D.PATHS["envipe"] / f"conjunto_de_datos_envipe_{year}_csv.zip")
    sexo = {k: D.norm(v) for k, v in _catalog(z, f"tper_vic1_envipe{year}/catalogos/sexo.csv").items()}
    D.check(sexo.get("1") == "HOMBRE" and sexo.get("2") == "MUJER", f"catálogo SEXO {year}: {sexo}")


def _pdiff(df, ya, da, yb, db) -> list:
    """(ya en da) − (yb en db) sobre la misma muestra, con variables y dominios distintos; IC95 simple."""
    ra, za = _lin(df, ya.astype(float) * da, da.astype(float))
    rb, zb = _lin(df, yb.astype(float) * db, db.astype(float))
    se = np.sqrt(_var(df, za - zb)); d = ra - rb
    return [round(d, 4), round(d - 1.96 * se, 4), round(d + 1.96 * se, 4)]


def _mean(df, y, dom) -> list:
    x = dom.astype(float)
    r, se = ratio_se(df, y * x, x)
    return [round(r, 1), round(r - 1.96 * se, 1), round(r + 1.96 * se, 1)]


def _total(df, y, dom) -> list:
    z = df.w * y * dom
    t = z.sum(); se = np.sqrt(_var(df, z))
    return [round(t), round(t - 1.96 * se), round(t + 1.96 * se)]


def edition(year: int) -> tuple[dict, dict]:
    df, _ = load_personas(year)
    _checks(df, year)
    m = D.municipios()
    reg = dict(zip(m.cvegeo.str[-3:], m.region))
    jal = df.ent == 14
    nat = df.ent.notna()
    amg = jal & (df.CVE_MUN.str.strip().str.zfill(3).map(reg) == "Área metropolitana de Guadalajara")
    resto = jal & ~amg
    doms = {"jalisco": jal, "amg": amg, "resto_del_estado": resto, "nacional": nat}
    mujer, hombre = _code(df, "SEXO") == "2", _code(df, "SEXO") == "1"

    act = {c: _code(df, f"AP4_10_{c}") for c in ACT}
    alguna = pd.concat([(a == "1") for a in act.values()], axis=1).any(axis=1)
    med = {c: _code(df, f"AP4_11_{c}") for c in MED}
    alguna_med = pd.concat([(a == "1") for a in med.values()], axis=1).any(axis=1)

    se = {}  # (p, se) para comparar ediciones y validar
    out = {"personas_muestra": {k: int(v.sum()) for k, v in doms.items()}, "actividades": [], "medidas": [],
           "problemas_colonia": [], "sensacion_de_ser_victima": []}
    for c, lab in ACT.items():
        ap = act[c] != "3"
        row = {"codigo": c, "actividad": lab}
        for k, d in doms.items():
            row[k] = prop(df, act[c] == "1", d & ap)
            se[("act", c, k)] = _se(df, act[c] == "1", d & ap)
        row["diferencia_jalisco_menos_nacional"] = diff(df, act[c] == "1", jal & ap, nat & ap)
        row["personas_muestra_jalisco_aplica"] = int((jal & ap).sum())
        out["actividades"].append(row)
    out["al_menos_una_actividad"] = {k: prop(df, alguna, d) for k, d in doms.items()}
    for k, d in doms.items():
        se[("alguna", k)] = _se(df, alguna, d)
    out["diferencia_amg_menos_resto_al_menos_una"] = diff(df, alguna, amg, resto)
    out["diferencia_jalisco_menos_nacional_al_menos_una"] = diff(df, alguna, jal, nat)
    noche = act["01"] == "1"; ap_noche = act["01"] != "3"
    out["salir_de_noche_por_sexo"] = {"mujeres": prop(df, noche, jal & mujer & ap_noche),
                                      "hombres": prop(df, noche, jal & hombre & ap_noche),
                                      "diferencia_mujeres_menos_hombres": diff(df, noche, jal & mujer & ap_noche, jal & hombre & ap_noche),
                                      "nacional_diferencia": diff(df, noche, nat & mujer & ap_noche, nat & hombre & ap_noche)}
    out["al_menos_una_por_sexo_jalisco"] = {"mujeres": prop(df, alguna, jal & mujer), "hombres": prop(df, alguna, jal & hombre),
                                             "diferencia": diff(df, alguna, jal & mujer, jal & hombre)}

    for c, lab in MED.items():
        out["medidas"].append({"codigo": c, "medida": lab, **{k: prop(df, med[c] == "1", d) for k, d in doms.items()}})
    out["alguna_medida"] = {k: prop(df, alguna_med, d) for k, d in doms.items()}
    out["diferencia_jalisco_menos_nacional_alguna_medida"] = diff(df, alguna_med, jal, nat)
    out["medidas_menos_actividades_jalisco"] = _paired(df, alguna_med, alguna, jal)

    g = _code(df, "AP4_12")
    num = pd.to_numeric(g.where(g.str.fullmatch(r"\d{1,7}")), errors="coerce")
    ns = num == 9999999
    monto = num.where(~ns)
    gasta = monto > 0
    # Flujo verificado: AP4_12 se pregunta sólo y siempre a quien tomó alguna medida; en 2026 también cuentan las
    # medidas electrónicas nuevas (AP4_11A_1..5), así que el gasto de 2026 no es comparable con el de 2025.
    filtro = alguna_med
    if year == 2026:
        filtro = alguna_med | pd.concat([_code(df, f"AP4_11A_{c}") == "1" for c in range(1, 6)], axis=1).any(axis=1)
    D.check((num.notna() == filtro).all(), f"AP4_12 {year}: no coincide con el filtro de medidas")
    out["gasto"] = {"pregunta_contestada_jalisco": int((jal & num.notna()).sum()),
                    "no_sabe_jalisco": int((jal & ns).sum()),
                    "incluye_medidas_electronicas": year == 2026}
    for k, d in doms.items():
        out["gasto"][k] = {"proporcion_que_gasta": prop(df, gasta.fillna(False), d & ~ns),
                           "gasto_medio_entre_quienes_gastan_pesos": _mean(df, monto.fillna(0), d & gasta.fillna(False)),
                           "casos_muestra_gastan": int((d & gasta.fillna(False)).sum())}
    out["gasto"]["jalisco"]["total_estatal_pesos"] = _total(df, monto.fillna(0), (jal & ~ns).astype(float))

    for c, lab in PROB.items():
        p8, p9 = _code(df, f"AP4_8_{c}"), _code(df, f"AP4_9_{c}")
        rep = p8 == "1"
        out["problemas_colonia"].append({"codigo": c, "problema": lab,
                                         **{f"reporta_{k}": prop(df, rep, d) for k, d in doms.items()},
                                         **{f"vecinos_se_organizan_{k}": prop(df, p9 == "1", d & rep) for k, d in doms.items()},
                                         "personas_muestra_jalisco_reportan": int((jal & rep).sum())})
    rob, alu = _code(df, "AP4_8_5") == "1", _code(df, "AP4_8_1") == "1"
    out["organizacion_robos_menos_alumbrado_jalisco"] = _pdiff(df, _code(df, "AP4_9_5") == "1", jal & rob,
                                                               _code(df, "AP4_9_1") == "1", jal & alu)
    for c, lab in SENS.items():
        v = _code(df, f"AP4_6_{c}")
        out["sensacion_de_ser_victima"].append({"codigo": c, "delito": lab,
                                                **{k: prop(df, v == "1", d & (v != "3")) for k, d in doms.items()}})
    if year == 2026:
        v = _code(df, "AP4_10_17")
        out["llegar_tarde_o_dejar_casa_sola_2026"] = {k: prop(df, v == "1", d & (v != "3")) for k, d in doms.items()}
        out["actividades_en_linea_2026"] = []
        for c, lab in EN_LINEA.items():
            v = _code(df, f"AP4_10A_{c}")
            out["actividades_en_linea_2026"].append({"codigo": c, "actividad": lab,
                                                     **{k: prop(df, v == "1", d & (v != "3")) for k, d in (("jalisco", jal), ("nacional", nat))}})
    return out, se


def _tab532(year: int, suf: str) -> dict:
    import openpyxl
    path = D.PATHS["envipe_tab"] / f"V_percepcion_seguridad_{year}_{suf}.xlsx"
    D._sha(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    hits = []
    for n in wb.sheetnames[1:-1]:
        t = _clean(" ".join(str(r[0]) for r in wb[n].iter_rows(min_row=2, max_row=6, values_only=True) if r and r[0]))
        if ("Población de 18 años y más por entidad federativa y actividad cotidiana según condición de haberla dejado" in t):
            hits.append(n)
    D.check(len(hits) == 1, f"tabulado {year} {suf}: cuadro de actividades encontrado {len(hits)} veces")
    rows = list(wb[hits[0]].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if any(_clean(c) == "Relativos" for c in r if c is not None))
    si = next(j for j, c in enumerate(rows[h - 1]) if c is not None and _clean(c) == "Sí")
    rel = next(j for j in range(si, len(rows[h])) if _clean(rows[h][j]) == "Relativos")
    labels = set(ACT.values())
    out, dom = {}, None
    for r in rows[h + 1:]:
        lab = _clean(r[0]) if r and r[0] is not None else ""
        if not lab:
            continue
        if lab.startswith(("Nota", "Fuente")):
            break
        if lab in labels:
            out[(dom, lab)] = float(r[rel])
        elif lab.startswith("Llegar muy tarde"):  # AP4_10_17, sólo 2026; no es encabezado de entidad
            continue
        else:
            dom = lab
    return out


def validacion(year: int, se: dict) -> dict:
    est, err = _tab532(year, "est"), _tab532(year, "err")
    rows = []
    for k, dname in (("jalisco", "Jalisco"), ("nacional", "Estados Unidos Mexicanos")):
        for c, lab in ACT.items():
            D.check((dname, lab) in est and (dname, lab) in err, f"tabulado {year}: falta {dname} / {lab}")
            p, s = se[("act", c, k)]
            rows.append({"dominio": k, "actividad": lab, "propio": round(p * 100, 2), "inegi": est[(dname, lab)],
                         "diferencia": round(p * 100 - est[(dname, lab)], 2),
                         "ee_propio": round(s * 100, 3), "ee_inegi": err[(dname, lab)],
                         "razon_ee": round(s * 100 / err[(dname, lab)], 3)})
    return {"cifras": rows, "max_diferencia": max(abs(r["diferencia"]) for r in rows),
            "razon_ee_min": min(r["razon_ee"] for r in rows), "razon_ee_max": max(r["razon_ee"] for r in rows)}


def run() -> dict:
    res, se = {}, {}
    for y in (2026, 2025):
        res[y], se[y] = edition(y)
    val = {y: validacion(y, se[y]) for y in (2026, 2025)}
    cambio = {k: _between(se[2026][("alguna", k)], se[2025][("alguna", k)]) for k in ("jalisco", "amg", "resto_del_estado", "nacional")}
    cambio_act = {ACT[c]: {k: _between(se[2026][("act", c, k)], se[2025][("act", c, k)])["razon"] for k in ("jalisco", "nacional")}
                  for c in ACT}
    j = res[2026]
    h = {"H22a": j["al_menos_una_actividad"]["jalisco"][1] > 0.5,
         "H22b": cambio["jalisco"]["razon"][1] > 0.90,
         "H22c": j["salir_de_noche_por_sexo"]["diferencia_mujeres_menos_hombres"][1] > 0,
         "H22d": j["diferencia_amg_menos_resto_al_menos_una"][1] > 0,
         "H22e": j["medidas_menos_actividades_jalisco"][2] < 0,
         "H22f": j["organizacion_robos_menos_alumbrado_jalisco"][2] < 0,
         "H22g": all(v["max_diferencia"] <= 0.5 for v in val.values()),
         "H22h": all(0.90 <= v["razon_ee_min"] and v["razon_ee_max"] <= 1.10 for v in val.values()),
         "E22i": "exploratoria"}
    return {"pregunta": "¿Qué deja de hacer la gente en Jalisco por miedo, y cómo se protege?",
            "ediciones": {"2026": "febrero-abril 2026", "2025": "marzo-abril 2025"},
            "resultados": {str(y): v for y, v in res.items()},
            "cambio_2025_2026_al_menos_una": cambio, "cambio_2025_2026_por_actividad_razon": cambio_act,
            "validacion_tabulado_5_32": {str(y): v for y, v in val.items()},
            "hipotesis": h}

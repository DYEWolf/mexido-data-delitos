"""Pieza 16 — ¿Qué tan inseguro se siente Jalisco, y la percepción sigue la baja del homicidio registrado?

Datos: ENVIPE 2026 (percepción medida en febrero-abril de 2026) y ENVIPE 2025 (marzo-abril de 2025), tabla de personas de
18 años y más (tper_vic1), módulo AP4. Variables verificadas en el diccionario y los catálogos de cada ZIP antes de
calcular (24-09-2026):
  - AP4_3_1 / AP4_3_2 / AP4_3_3: seguridad en su colonia o localidad / municipio / estado (1 seguro, 2 inseguro,
    9 no sabe / no responde); se preguntan a todas las personas elegidas.
  - AP4_4_01..12: seguridad en 12 lugares (1 seguro, 2 inseguro, 3 no aplica, 9 no sabe / no responde).
  - AP4_5_xx: incivilidades en su colonia (1 = sí; 0 = no se declaró); los códigos 01-17 tienen la misma etiqueta en
    los dos años; en 2026 cambian 18-20.
  - AP4_7_2: tendencia esperada de la seguridad en su estado (1 mejorará ... 4 empeorará, 9 no sabe).
  - Entidad y municipio de residencia (CVE_ENT, CVE_MUN), FAC_ELE, EST_DIS, UPM_DIS.
Ningún resultado de este módulo se había visto antes de escribir estas hipótesis; sí se conocen la baja del homicidio
registrado en 2025 (pieza 12: 0.66 [0.60-0.71] en carpetas) y el tamaño de muestra por región (AMG 1,754 personas en
2026; resto del estado, incluidos los Altos, 910).

Hipótesis (declaradas el 24-09-2026 antes de calcular):
  H16a. La mayoría de los adultos de Jalisco se siente insegura en su estado (ENVIPE 2026): IC95 inferior > 50%.
        Refutación: IC95 inferior <= 50%.
  H16b. Jalisco no se aparta del país en la percepción de inseguridad en el estado: IC95 de la diferencia Jalisco −
        nacional dentro de ±5 puntos. Refutación: el IC95 sale de ese rango.
  H16c. En Jalisco la inseguridad se percibe más lejos que cerca: estado > municipio > colonia, y cada diferencia
        (pareada, mismas personas) tiene IC95 inferior > 0. Refutación: alguna de las dos no lo cumple.
  H16d. La percepción no sigue la baja del homicidio registrado: la razón 2026/2025 de la proporción de adultos de
        Jalisco que se siente insegura en su estado tiene IC95 inferior > 0.90 (se descarta una baja mayor a 10%,
        frente a −34% del homicidio). Refutación: IC95 inferior <= 0.90.
  H16e. En el área metropolitana (10 municipios, como en las piezas 1, 5, 12 y 14) la inseguridad en la colonia es
        mayor que en el resto del estado: IC95 de la diferencia AMG − resto > 0. Refutación: IC95 inferior <= 0.
  H16f. El cajero automático en la vía pública es el lugar donde más adultos de Jalisco se sienten inseguros
        (entre quienes el lugar les aplica). Refutación: otro lugar tiene una estimación puntual mayor.

Robustez y descripción (sin hipótesis): todo con ENVIPE 2025; H16d también en colonia y municipio, y como diferencia
de cambios Jalisco − nacional; lugares, incivilidades 01-17 y tendencia esperada, Jalisco contra nacional.

Exploratorio (añadido el 24-09-2026 después de ver los resultados de las dos ediciones):
  E16g. Cambio 2025 → 2026 en Jalisco y en el país de las 17 incivilidades comparables (todas, no solo homicidios y
        disparos, que fueron las que motivaron la pregunta).

Validación contra los tabulados publicados por INEGI (declarada el 25-09-2026, antes de leer cualquier cifra de los
tabulados). Insumo: tabulados básicos "V. Percepción sobre la seguridad pública" de la ENVIPE 2025 y 2026 (estimaciones
y errores estándar; carpeta s3-envipe-tabulados-*). Antes de declarar solo se revisó la estructura: índice, títulos,
encabezados, etiquetas de fila y notas al pie. Según las notas, en los cuadros de colonia, municipio, entidad y
tendencia el total incluye los casos "no especificado" (igual que aquí). En el cuadro de espacios se leyó que se
excluyen "no sabe / no responde" y los lugares que no aplican. Corrección (25-09-2026, después de ver las cifras): esa
nota se refiere a las columnas que muestra el cuadro. El denominador de INEGI incluye "no sabe / no responde", igual que
el cálculo de lugares de este script, y se comprobó con los datos (escuela, transporte público y parque en 2025). La
comparación de lugares usa el cálculo original; no afecta ninguna hipótesis.
  H16h. Las estimaciones de este script reproducen el porcentaje "inseguro" publicado para colonia, municipio y entidad,
        en Jalisco y en el total nacional, en las dos ediciones (12 cifras): diferencia absoluta <= 0.5 puntos en todas.
        Refutación: alguna difiere más de 0.5 puntos.
  H16i. Los errores estándar de este script para esas 12 cifras están dentro de ±10% de los publicados por INEGI
        (razón propio/INEGI entre 0.90 y 1.10). Refutación: alguno fuera de ese rango.
  Descripción (sin hipótesis): lugares con el alcance de INEGI y tendencia esperada en la entidad (4 categorías),
  diferencia máxima con lo publicado. Los intervalos publicados son al 90%; no se comparan.

Estimador: el de las piezas 9 y 13 (razón ponderada con FAC_ELE, linealización de Taylor con EST_DIS y UPM_DIS, dominios
sin eliminar UPM del diseño, IC95 en escala logit). "No sabe / no responde" queda en el denominador. Diferencias dentro
de una misma edición: linealización sobre la muestra completa (IC95 simple). Entre ediciones (muestras
independientes): varianzas sumadas; la razón usa el método delta en escala logarítmica.
"""
from __future__ import annotations

import re

import numpy as np
import pandas as pd

from mvj import data as D
from pieces.p9_cifra_negra import ratio_se
from pieces.p13_no_denuncia_confianza import _lin, _var, diff, load_personas, prop

NIVELES = {"AP4_3_1": "colonia o localidad", "AP4_3_2": "municipio", "AP4_3_3": "estado"}
LUGARES = {"01": "Casa", "02": "Trabajo", "03": "Calle", "04": "Escuela", "05": "Mercado", "06": "Centro comercial",
           "07": "Banco", "08": "Cajero automático en la vía pública", "09": "Transporte público", "10": "Automóvil",
           "11": "Carretera", "12": "Parque o centro recreativo"}
INCIVILIDADES = {"01": "Consumo de alcohol en la calle", "02": "Pandillerismo o bandas violentas", "03": "Riñas entre vecinos",
                 "04": "Venta ilegal de alcohol", "05": "Venta de productos piratas", "06": "Violencia policiaca contra ciudadanos",
                 "07": "Invasión de predios", "08": "Consumo de droga", "09": "Robos o asaltos frecuentes", "10": "Venta de droga",
                 "11": "Disparos frecuentes", "12": "Prostitución", "13": "Secuestros", "14": "Homicidios",
                 "15": "Extorsiones (cobro de piso)", "16": "Robo o venta ilegal de combustible (huachicol)",
                 "17": "Tomas irregulares de luz"}


def _check(df: pd.DataFrame, year: int) -> None:
    for v in NIVELES:
        D.check(set(df[v].fillna("NA").str.strip()) <= {"1", "2", "9"}, f"{v} {year}: códigos fuera de {{1, 2, 9}} o vacíos")
    for c in LUGARES:
        D.check(set(df[f"AP4_4_{c}"].fillna("NA").str.strip()) <= {"1", "2", "3", "9"}, f"AP4_4_{c} {year}: códigos fuera de catálogo")
    for c in INCIVILIDADES:
        D.check(set(df[f"AP4_5_{c}"].fillna("NA").str.strip()) <= {"0", "1"}, f"AP4_5_{c} {year}: códigos fuera de catálogo")
    D.check(set(df.AP4_7_2.fillna("NA").str.strip()) <= {"1", "2", "3", "4", "9"}, f"AP4_7_2 {year}: códigos fuera de catálogo")
    D.check(df.CVE_MUN.notna().all(), f"CVE_MUN {year}: sin vacíos")


def _catalogs(year: int) -> None:
    import zipfile
    from pieces.p13_no_denuncia_confianza import _catalog
    z = zipfile.ZipFile(D.PATHS["envipe"] / f"conjunto_de_datos_envipe_{year}_csv.zip")
    base = f"tper_vic1_envipe{year}/catalogos"
    for v in NIVELES:
        c = {k: D.norm(x) for k, x in _catalog(z, f"{base}/{v.lower()}.csv").items()}
        D.check(c.get("1") == "SEGURO" and c.get("2") == "INSEGURO", f"catálogo {v} {year}: {c}")
    for k in LUGARES:
        c = {k2: D.norm(x) for k2, x in _catalog(z, f"{base}/ap4_4_{k}.csv").items()}
        D.check(c.get("2", "").startswith("INSEGURO") and c.get("3") == "NO APLICA", f"catálogo AP4_4_{k} {year}")
    c = {k: D.norm(x) for k, x in _catalog(z, f"{base}/ap4_7_2.csv").items()}
    D.check(c.get("4") == "EMPEORARA", f"catálogo AP4_7_2 {year}")


def _se(df, y, dom) -> tuple[float, float]:
    x = dom.astype(float)
    return ratio_se(df, y.astype(float) * x, x)


def _paired(df, a, b, dom) -> list:
    x = dom.astype(float)
    r, z = _lin(df, (a.astype(float) - b.astype(float)) * x, x)
    se = np.sqrt(_var(df, z))
    return [round(r, 4), round(r - 1.96 * se, 4), round(r + 1.96 * se, 4)]


def edition(year: int) -> tuple[dict, dict]:
    df, _ = load_personas(year)
    _catalogs(year); _check(df, year)
    m = D.municipios()
    reg = dict(zip(m.cvegeo.str[-3:], m.region))
    jal = df.ent == 14
    nat = df.ent.notna()
    D.check(set(df.loc[jal, "CVE_MUN"].str.strip().str.zfill(3)) <= set(reg), f"CVE_MUN Jalisco {year} fuera del catálogo INEGI")
    amg = jal & (df.CVE_MUN.str.strip().str.zfill(3).map(reg) == "Área metropolitana de Guadalajara")
    resto = jal & ~amg
    ins = {v: df[v].str.strip() == "2" for v in NIVELES}
    out = {"personas_muestra": {"jalisco": int(jal.sum()), "amg": int(amg.sum()), "resto_del_estado": int(resto.sum()),
                                "nacional": int(nat.sum())},
           "inseguro": {}, "gradiente_jalisco": {}, "lugares": [], "incivilidades": [], "tendencia_estado": {}}
    se = {}
    for v, lab in NIVELES.items():
        out["inseguro"][lab] = {"jalisco": prop(df, ins[v], jal), "nacional": prop(df, ins[v], nat),
                                "diferencia_jalisco_menos_nacional": diff(df, ins[v], jal, nat),
                                "amg": prop(df, ins[v], amg), "resto_del_estado": prop(df, ins[v], resto),
                                "diferencia_amg_menos_resto": diff(df, ins[v], amg, resto)}
        se[lab] = {"jalisco": _se(df, ins[v], jal), "nacional": _se(df, ins[v], nat)}
    out["gradiente_jalisco"] = {"estado_menos_municipio": _paired(df, ins["AP4_3_3"], ins["AP4_3_2"], jal),
                                "municipio_menos_colonia": _paired(df, ins["AP4_3_2"], ins["AP4_3_1"], jal)}
    for c, lab in LUGARES.items():
        v = df[f"AP4_4_{c}"].str.strip()
        applies = v != "3"
        out["lugares"].append({"codigo": c, "lugar": lab, "jalisco": prop(df, v == "2", jal & applies),
                               "nacional": prop(df, v == "2", nat & applies),
                               "diferencia_jalisco_menos_nacional": diff(df, v == "2", jal & applies, nat & applies),
                               "personas_muestra_jalisco": int((jal & applies).sum())})
    out["lugares"].sort(key=lambda r: -r["jalisco"][0])
    for c, lab in INCIVILIDADES.items():
        y = df[f"AP4_5_{c}"].str.strip() == "1"
        out["incivilidades"].append({"codigo": c, "incivilidad": lab, "jalisco": prop(df, y, jal), "nacional": prop(df, y, nat),
                                     "diferencia_jalisco_menos_nacional": diff(df, y, jal, nat),
                                     "amg": prop(df, y, amg), "resto_del_estado": prop(df, y, resto)})
        se[f"inc_{c}"] = {"jalisco": _se(df, y, jal), "nacional": _se(df, y, nat)}
    out["incivilidades"].sort(key=lambda r: -r["jalisco"][0])
    t = df.AP4_7_2.str.strip()
    for code, lab in (("1", "mejorará"), ("2", "seguirá igual de bien"), ("3", "seguirá igual de mal"), ("4", "empeorará")):
        out["tendencia_estado"][lab] = {"jalisco": prop(df, t == code, jal), "nacional": prop(df, t == code, nat)}
    return out, se


def _between(a: tuple, b: tuple) -> dict:
    """a = (p, se) newer edition, b = older; independent samples."""
    (pa, sa), (pb, sb) = a, b
    d, sd = pa - pb, np.sqrt(sa ** 2 + sb ** 2)
    r, sl = pa / pb, np.sqrt((sa / pa) ** 2 + (sb / pb) ** 2)
    return {"diferencia": [round(d, 4), round(d - 1.96 * sd, 4), round(d + 1.96 * sd, 4)],
            "razon": [round(r, 3), round(float(r * np.exp(-1.96 * sl)), 3), round(float(r * np.exp(1.96 * sl)), 3)]}


TAB = {"colonia o localidad": "según percepción sobre la seguridad en colonia o localidad",
       "municipio": "según percepción sobre la seguridad en municipio o demarcación territorial",
       "estado": "según percepción sobre la seguridad en entidad federativa",
       "espacios": "y espacio público o privado según percepción de seguridad en éstos",
       "tendencia": "según percepción sobre la tendencia de la seguridad pública en su entidad federativa"}
INEGI_LUGARES = {"01": "Su casa", "02": "Su trabajo", "03": "La calle", "04": "La escuela", "05": "El mercado",
                 "06": "El centro comercial", "07": "El banco", "08": "El cajero automático en la vía pública",
                 "09": "El transporte público", "10": "El automóvil", "11": "La carretera", "12": "El parque o centro recreativo"}
INEGI_TENDENCIA = {"mejorará": "Mejorará", "seguirá igual de bien": "Seguirá igual de bien",
                   "seguirá igual de mal": "Seguirá igual de mal", "empeorará": "Empeorará"}
DOMS = {"jalisco": "Jalisco", "nacional": "Estados Unidos Mexicanos"}


def _clean(x) -> str:
    return re.sub(r"\s+", " ", str(x)).strip()


def _tabulado(year: int, suf: str) -> dict:
    """Cuadros de percepción del tabulado V de INEGI: {cuadro: {(dominio, fila): {categoría: relativo}}}."""
    import openpyxl
    path = D.PATHS["envipe_tab"] / f"V_percepcion_seguridad_{year}_{suf}.xlsx"
    D._sha(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for key, title in TAB.items():
        hits = []
        for n in wb.sheetnames[1:-1]:
            t = _clean(" ".join(str(r[0]) for r in wb[n].iter_rows(min_row=2, max_row=6, values_only=True) if r and r[0]))
            if "Población de 18 años y más por entidad federativa" in t and title in t and not re.search(r"sexo|grupos de edad|escolaridad", t):
                hits.append(n)
        D.check(len(hits) == 1, f"tabulado {year} {suf}: cuadro '{key}' encontrado {len(hits)} veces")
        rows = list(wb[hits[0]].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows) if any(_clean(c) == "Relativos" for c in r if c is not None))
        cols, cur = {}, None
        for j, (c, r) in enumerate(zip(rows[h - 1], rows[h])):
            if c is not None and _clean(c):
                cur = _clean(c).replace("(a)", "")
            if r is not None and _clean(r) == "Relativos":
                cols[cur] = j
        table, dom = {}, None
        for r in rows[h + 1:]:
            lab = _clean(r[0]) if r and r[0] is not None else ""
            if not lab:
                continue
            if lab.startswith(("Nota", "Nivel de precisión", "Fuente")):
                break
            vals = {k: r[j] for k, j in cols.items()}
            if all(v is None for v in vals.values()):
                dom = lab  # encabezado de entidad en el cuadro de espacios
                continue
            if key == "espacios" and lab in INEGI_LUGARES.values():
                table[(dom, lab)] = {k: float(v) for k, v in vals.items()}
            else:
                dom = lab
                table[(lab, None)] = {k: float(v) for k, v in vals.items()}
        for d in DOMS.values():
            D.check(any(k[0] == d for k in table), f"tabulado {year} {suf} '{key}': falta {d}")
        out[key] = table
    return out


def validacion(year: int, res: dict, se: dict) -> dict:
    est, err = _tabulado(year, "est"), _tabulado(year, "err")
    niveles = []
    for lab in NIVELES.values():
        for dom, name in DOMS.items():
            e, s_ = est[lab][(name, None)], err[lab][(name, None)]
            D.check(90 <= e["Seguro"] + e["Inseguro"] <= 100.05, f"tabulado {year} {lab} {name}: seguro + inseguro fuera de [90, 100]")
            p, sp = se[lab][dom]
            niveles.append({"nivel": lab, "dominio": dom, "propio": round(p * 100, 2), "inegi": round(e["Inseguro"], 3),
                            "diferencia_puntos": round(p * 100 - e["Inseguro"], 2),
                            "ee_propio": round(sp * 100, 3), "ee_inegi": round(s_["Inseguro"], 3), "razon_ee": round(sp * 100 / s_["Inseguro"], 3)})
    lugares = []
    for c, name in INEGI_LUGARES.items():
        for dom, dname in DOMS.items():
            e = est["espacios"][(dname, name)]
            ours = next(r for r in res["lugares"] if r["codigo"] == c)[dom][0] * 100
            lugares.append({"lugar": name, "dominio": dom, "propio": round(ours, 2), "inegi": round(e["Inseguro"], 3),
                            "diferencia_puntos": round(ours - e["Inseguro"], 2)})
    tendencia = []
    for lab, name in INEGI_TENDENCIA.items():
        for dom, dname in DOMS.items():
            ours = res["tendencia_estado"][lab][dom][0] * 100
            e = est["tendencia"][(dname, None)][name]
            tendencia.append({"categoria": name, "dominio": dom, "propio": round(ours, 2), "inegi": round(e, 3), "diferencia_puntos": round(ours - e, 2)})
    return {"niveles": niveles, "lugares": lugares, "tendencia": tendencia,
            "max_dif_lugares": max(abs(r["diferencia_puntos"]) for r in lugares),
            "max_dif_tendencia": max(abs(r["diferencia_puntos"]) for r in tendencia)}


def run() -> dict:
    r26, se26 = edition(2026)
    r25, se25 = edition(2025)
    cambio = {}
    for lab in NIVELES.values():
        j, n = _between(se26[lab]["jalisco"], se25[lab]["jalisco"]), _between(se26[lab]["nacional"], se25[lab]["nacional"])
        (dj, sj), (dn, sn) = [(x["diferencia"][0], (x["diferencia"][2] - x["diferencia"][0]) / 1.96) for x in (j, n)]
        # Jalisco is inside the national sample; treating the two changes as independent slightly overstates the SE.
        dd, sdd = dj - dn, np.sqrt(sj ** 2 + sn ** 2)
        cambio[lab] = {"jalisco_2026_vs_2025": j, "nacional_2026_vs_2025": n,
                       "cambio_jalisco_menos_cambio_nacional": [round(dd, 4), round(dd - 1.96 * sdd, 4), round(dd + 1.96 * sdd, 4)]}
    val = {2026: validacion(2026, r26, se26), 2025: validacion(2025, r25, se25)}
    niv = [r for v in val.values() for r in v["niveles"]]
    e16g = []
    for c, lab in INCIVILIDADES.items():
        k = f"inc_{c}"
        e16g.append({"codigo": c, "incivilidad": lab,
                     "jalisco_2026_vs_2025": _between(se26[k]["jalisco"], se25[k]["jalisco"]),
                     "nacional_2026_vs_2025": _between(se26[k]["nacional"], se25[k]["nacional"])})
    e16g.sort(key=lambda r: r["jalisco_2026_vs_2025"]["diferencia"][0])
    est = r26["inseguro"]["estado"]
    g = r26["gradiente_jalisco"]
    col = r26["inseguro"]["colonia o localidad"]
    return {
        "pregunta": "¿Qué tan inseguro se siente Jalisco, y la percepción sigue la baja del homicidio registrado?",
        "resultados": {2026: {"levantamiento": "febrero-abril 2026", **r26}, 2025: {"levantamiento": "marzo-abril 2025", **r25}},
        "cambio_2025_a_2026": cambio,
        "E16g_cambio_incivilidades": e16g,
        "validacion_tabulados_inegi": val,
        "hipotesis": {
            "H16a": est["jalisco"][1] > 0.50,
            "H16b": -0.05 <= est["diferencia_jalisco_menos_nacional"][1] and est["diferencia_jalisco_menos_nacional"][2] <= 0.05,
            "H16c": g["estado_menos_municipio"][1] > 0 and g["municipio_menos_colonia"][1] > 0,
            "H16d": cambio["estado"]["jalisco_2026_vs_2025"]["razon"][1] > 0.90,
            "H16e": col["diferencia_amg_menos_resto"][1] > 0,
            "H16f": r26["lugares"][0]["codigo"] == "08",
            "E16g": "exploratoria",
            "H16h": all(abs(r["diferencia_puntos"]) <= 0.5 for r in niv),
            "H16i": all(0.90 <= r["razon_ee"] <= 1.10 for r in niv),
        },
    }
